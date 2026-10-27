from datetime import date, datetime, timezone as dt_timezone
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from requests.exceptions import Timeout

from dev.models import Dev
from sprint.scope_flow import build_scope_flow
from sprint.logged_report import (
    HEADERS,
    UNASSIGNED_HEADERS,
    DevIdentity,
    LoggedWorkbookBuilder,
    build_dev_rows,
    build_unassigned_rows,
    logged_hours,
    person_matches_dev,
    resolve_dev,
    sheet_title,
    sprint_window,
    status_at,
)
from sprint.models import SprintPeriod
from team.models import Team, TeamDev

from JiraFox.services.dashboard import (
    build_dashboard_bottlenecks,
    build_dashboard_items,
    build_dashboard_kpis,
)
from sprint.views import build_weighted_throughput_distribution


class DashboardServicesTests(SimpleTestCase):
    def test_build_dashboard_kpis_handles_empty_and_populated_metrics(self):
        issues = [
            {
                "state": "In Progress",
                "kanban_metrics": {
                    "current_wip_age_days": 12,
                    "cycle_time_days": 6,
                    "lead_time_days": 8,
                },
            },
            {
                "state": "Done",
                "kanban_metrics": {
                    "current_wip_age_days": 2,
                    "cycle_time_days": 4,
                    "lead_time_days": 10,
                },
            },
        ]

        result = build_dashboard_kpis(issues, throughput_7d=5)

        self.assertEqual(result["wip"], 1)
        self.assertEqual(result["aging_wip"], 1)
        self.assertEqual(result["throughput_7d"], 5)
        self.assertEqual(result["cycle_p50_days"], 5.0)
        self.assertEqual(result["lead_p50_days"], 9.0)
        self.assertEqual(result["flow_eff_pct"], 55.6)
        self.assertEqual(result["aging_wip_pct"], 100.0)

    def test_build_dashboard_bottlenecks_sorts_by_average_age(self):
        issues = [
            {"state": "Review", "kanban_metrics": {"current_wip_age_days": 3}},
            {"state": "Review", "kanban_metrics": {"current_wip_age_days": 9}},
            {"state": "Testing", "kanban_metrics": {"current_wip_age_days": 4}},
        ]

        result = build_dashboard_bottlenecks(issues)

        self.assertEqual(result[0]["status"], "Review")
        self.assertEqual(result[0]["avgAge"], 6.0)
        self.assertEqual(result[0]["maxAge"], 9)

    def test_build_dashboard_items_maps_expected_fields(self):
        issues = [
            {
                "key": "ABC-1",
                "summary": "Titre",
                "state": "In Progress",
                "assignee": {"displayName": "Jane Doe"},
                "url": "https://jira.local/browse/ABC-1",
                "kanban_metrics": {"current_wip_age_days": 7},
            }
        ]

        result = build_dashboard_items(issues)

        self.assertEqual(result[0]["key"], "ABC-1")
        self.assertEqual(result[0]["title"], "Titre")
        self.assertEqual(result[0]["status"], "In Progress")
        self.assertEqual(result[0]["ageDays"], 7)
        self.assertEqual(result[0]["assignee"], "Jane Doe")
        self.assertEqual(result[0]["detail_url"], "/issue/ABC-1/")


class SprintPagesTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="alice", password="secret")
        self.client.force_login(self.user)

    @patch("sprint.views.get_connected_jira_for_user")
    def test_sprint_list_page_renders(self, mock_get_jira):
        mock_get_jira.return_value = SimpleNamespace(
            cfg=SimpleNamespace(jira_board_id=42),
            list_sprints_for_current_board=lambda states: [
                {
                    "id": 12,
                    "name": "Sprint 12",
                    "state": "active",
                    "startDate": "2026-03-01T08:00:00.000+0000",
                    "endDate": "2026-03-14T18:00:00.000+0000",
                    "url": "https://jira.example/sprint/12",
                }
            ],
        )

        response = self.client.get(reverse("sprint-list"))

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "sprint/sprint_list.html")
        self.assertContains(response, "Sprint 12")

    @patch("sprint.views.get_connected_jira_for_user")
    def test_sprint_issues_api_returns_mapped_items(self, mock_get_jira):
        mock_get_jira.return_value = SimpleNamespace(
            JIRA_ONLY_STANDARD_TYPES="issuetype in standardIssueTypes()",
            get_cached_sprint_issues=lambda sprint_id, jql_extra=None, full=None: [
                {
                    "key": "ABC-1",
                    "summary": "Issue title",
                    "issueType": "Bug",
                    "state": "In Progress",
                    "assignee": {"displayName": "Jane Doe"},
                    "url": "https://jira.example/browse/ABC-1",
                }
            ],
        )

        response = self.client.get(reverse("sprint-api-issues", args=[12]))

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["items"][0]["key"], "ABC-1")
        self.assertEqual(payload["items"][0]["issueType"], "Bug")

    @patch("sprint.views.get_connected_jira_for_user")
    def test_sprint_wia_api_returns_statuses_and_items(self, mock_get_jira):
        mock_get_jira.return_value = SimpleNamespace(
            JIRA_ONLY_STANDARD_TYPES="issuetype in standardIssueTypes()",
            get_cached_sprint_issues=lambda sprint_id, jql_extra=None, full=None: [
                {
                    "key": "ABC-1",
                    "summary": "Issue title",
                    "state": "In Progress",
                    "url": "https://jira.example/browse/ABC-1",
                    "kanban_metrics": {
                        "current_wip_age_hours": 16,
                        "cycle_time_hours": 8,
                        "first_in_progress": "2026-03-01T08:00:00.000+0000",
                        "first_in_progress_display": "01/03/2026",
                    },
                }
            ],
        )

        response = self.client.get(reverse("sprint-api-wia", args=[12]))

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["statuses"], ["In Progress"])
        self.assertEqual(payload["items"][0]["key"], "ABC-1")


SPRINT = {
    "id": 2446,
    "startDate": "2026-03-02T08:00:00.000+01:00",
    "endDate": "2026-03-13T18:00:00.000+01:00",
    "completeDate": None,
}
ALICE = {"displayName": "Martin, Alice", "name": "AMARTIN"}
BOB = {"displayName": "Dupont, Bob", "name": "BDUPONT"}
CARL = {"displayName": "Extern, Carl", "name": "CEXTERN"}
ISSUES = [
    {
        "id": "101",
        "key": "ABC-1",
        "summary": "Login page",
        "state": "Done",
        "created": "2026-02-20T10:00:00.000+0100",
        "assignee": ALICE,
        "original_estimate_hours": 8.0,
        "status_changes": [
            {"at": "2026-02-25T10:00:00.000+0100", "from": "Open", "to": "Ready"},
            {"at": "2026-03-03T09:00:00.000+0100", "from": "Ready", "to": "In Progress"},
            {"at": "2026-03-10T09:00:00.000+0100", "from": "In Progress", "to": "Done"},
        ],
        "worklogs": [
            {"author": ALICE, "started": "2026-03-03T09:00:00.000+0100", "hours": 3.0},
            {"author": ALICE, "started": "2026-03-04T09:00:00.000+0100", "hours": 2.5},
            {"author": ALICE, "started": "2026-02-27T09:00:00.000+0100", "hours": 4.0},  # before the sprint
            {"author": BOB, "started": "2026-03-05T09:00:00.000+0100", "hours": 1.0},
            {"author": CARL, "started": "2026-03-06T09:00:00.000+0100", "hours": 1.5},
            {"author": CARL, "started": "2026-03-09T09:00:00.000+0100", "hours": 0.5},
            {"author": CARL, "started": "2026-03-16T09:00:00.000+0100", "hours": 9.0},  # after the sprint
        ],
        "url": "https://jira.example/browse/ABC-1",
    },
    {
        "id": "102",
        "key": "ABC-2",
        "summary": "Hotfix",
        "state": "In Progress",
        "created": "2026-03-05T10:00:00.000+0100",
        "assignee": BOB,
        "original_estimate_hours": None,
        "status_changes": [],
        "worklogs": [],
        "url": "https://jira.example/browse/ABC-2",
    },
]
NOW = datetime(2026, 3, 20, 12, 0, tzinfo=dt_timezone.utc)


class LoggedReportTests(SimpleTestCase):
    def test_status_at_rebuilds_status_from_changelog(self):
        start, end = sprint_window(SPRINT, NOW)

        self.assertEqual(status_at(ISSUES[0], start), "Ready")
        self.assertEqual(status_at(ISSUES[0], end), "Done")
        self.assertIsNone(status_at(ISSUES[1], start))  # created during the sprint
        self.assertEqual(status_at(ISSUES[1], end), "In Progress")

    def test_logged_hours_only_counts_the_developer_within_the_sprint(self):
        start, end = sprint_window(SPRINT, NOW)

        self.assertEqual(logged_hours(ISSUES[0], "Alice", start, end), 5.5)
        self.assertEqual(logged_hours(ISSUES[0], "Bob", start, end), 1.0)

    def test_person_matching_accepts_first_name_display_name_and_username(self):
        self.assertTrue(person_matches_dev(ALICE, "alice"))
        self.assertTrue(person_matches_dev(ALICE, "Martin, Alice"))
        self.assertTrue(person_matches_dev(ALICE, "AMARTIN"))
        self.assertFalse(person_matches_dev(ALICE, "Bob"))

    def test_workbook_has_one_tab_per_dev_with_totals(self):
        content = LoggedWorkbookBuilder(sprint=SPRINT, issues=ISSUES, devs=["Alice", "Bob"], now=NOW).build()
        workbook = load_workbook(BytesIO(content))

        self.assertEqual(workbook.sheetnames, ["Alice", "Bob", "Unassigned", "Period"])
        alice = [[cell.value for cell in row] for row in workbook["Alice"].iter_rows()]
        self.assertEqual(alice[0], list(HEADERS))
        self.assertEqual(alice[1], ["ABC-1 - Login page", 8.0, "Ready", 5.5, "Done", "Yes"])
        self.assertEqual(alice[2], ["Total", "=SUM(B2:B2)", None, "=SUM(D2:D2)", None, None])
        self.assertEqual(alice[3], ["Outside sprint", None, None, '=SUMIF(F2:F2,"No",D2:D2)', None, None])

        bob = [[cell.value for cell in row] for row in workbook["Bob"].iter_rows()]
        # ABC-1 because Bob logged on it, ABC-2 because it is assigned to him.
        self.assertEqual(bob[1], ["ABC-1 - Login page", 8.0, "Ready", 1.0, "Done", "Yes"])
        self.assertEqual(bob[2], ["ABC-2 - Hotfix", 0.0, "Created during sprint", 0.0, "In Progress", "Yes"])
        self.assertEqual(bob[3][:4], ["Total", "=SUM(B2:B3)", None, "=SUM(D2:D3)"])

        unassigned = [[cell.value for cell in row] for row in workbook["Unassigned"].iter_rows()]
        self.assertEqual(unassigned[0], list(UNASSIGNED_HEADERS))
        self.assertEqual(
            unassigned[1], ["ABC-1 - Login page", 8.0, "Ready", 2.0, "Done", "Extern, Carl (CEXTERN)"]
        )
        self.assertEqual(unassigned[2][:4], ["Total", "=SUM(B2:B2)", None, "=SUM(D2:D2)"])

    def test_unassigned_rows_group_by_ticket_and_author(self):
        start, end = sprint_window(SPRINT, NOW)

        rows = build_unassigned_rows(ISSUES, ["Alice"], start, end)

        self.assertEqual(
            [(row["ticket"], row["author"], row["planned_hours"], row["logged_hours"]) for row in rows],
            [
                ("ABC-1 - Login page", "Dupont, Bob (BDUPONT)", 8.0, 1.0),
                # The ticket's planned time is only counted on its first row.
                ("ABC-1 - Login page", "Extern, Carl (CEXTERN)", None, 2.0),
            ],
        )

    def test_sheet_titles_are_valid_and_unique(self):
        used: set[str] = set()

        self.assertEqual(sheet_title("A/B", used), "A-B")
        self.assertEqual(sheet_title("a/b", used), "a-b (2)")
        self.assertLessEqual(len(sheet_title("x" * 40, used)), 31)

    def test_logged_hours_are_summed_before_rounding(self):
        start, end = sprint_window(SPRINT, NOW)
        issue = {
            **ISSUES[1],
            "worklogs": [
                {"author": ALICE, "started": f"2026-03-0{day}T09:00:00.000+0100", "hours": 0.33, "seconds": 1200}
                for day in (3, 4, 5)
            ],
        }

        self.assertEqual(logged_hours(issue, "Alice", start, end), 1.0)

    def test_sprint_not_started_has_no_logged_time(self):
        sprint = {**SPRINT, "startDate": None}
        start, end = sprint_window(sprint, NOW)

        self.assertEqual(logged_hours(ISSUES[0], "Alice", start, end), 0)
        self.assertEqual(build_unassigned_rows(ISSUES, ["Alice"], start, end), [])

    def test_each_worklog_is_attributed_to_one_developer_only(self):
        alice_bis = {"displayName": "Durand, Alice", "name": "ADURAND"}
        self.assertEqual(resolve_dev(ALICE, ["Alice", "AMARTIN"]), DevIdentity("AMARTIN"))  # exact match wins
        self.assertIsNone(resolve_dev(alice_bis, ["Alice", "alice"]))  # ambiguous first name
        self.assertEqual(resolve_dev(alice_bis, ["Alice", "Bob"]), DevIdentity("Alice"))
        self.assertIsNone(resolve_dev(CARL, ["Alice", "Bob"]))

        start, end = sprint_window(SPRINT, NOW)
        dev_names = ["Alice", "AMARTIN"]
        self.assertEqual(logged_hours(ISSUES[0], "AMARTIN", start, end, dev_names), 5.5)
        self.assertEqual(logged_hours(ISSUES[0], "Alice", start, end, dev_names), 0)

    def test_jira_username_matches_exactly_and_beats_names(self):
        alice = DevIdentity("Alice", jira_username="amartin")
        alice_bis = {"displayName": "Durand, Alice", "name": "ADURAND"}

        self.assertTrue(person_matches_dev(ALICE, alice))
        # With a Jira ID, a namesake is no longer taken for the developer...
        self.assertFalse(person_matches_dev(alice_bis, alice))
        self.assertIsNone(resolve_dev(alice_bis, [alice, "Bob"]))
        # ...and the ID wins over another developer matching by name.
        self.assertEqual(resolve_dev(ALICE, [alice, "Martin, Alice"]), alice)

    def test_time_logged_outside_the_sprint_is_listed_without_planned_time(self):
        start, end = sprint_window(SPRINT, NOW)
        support = {
            **ISSUES[1],
            "key": "SUP-9",
            "summary": "Support",
            "original_estimate_hours": 5.0,
            "assignee": BOB,
            "worklogs": [
                {"author": ALICE, "started": "2026-03-04T14:00:00.000+0100", "seconds": 5400},
                {"author": ALICE, "started": "2026-03-20T10:00:00.000+0100", "seconds": 3600},  # after the sprint
            ],
        }

        rows = build_dev_rows(ISSUES, "Alice", start, end, ["Alice", "Bob"], [support])
        self.assertEqual(
            [(row["ticket"], row["planned_hours"], row["logged_hours"], row["in_sprint"]) for row in rows],
            [("ABC-1 - Login page", 8.0, 5.5, True), ("SUP-9 - Support", None, 1.5, False)],
        )
        # Bob is the assignee but logged nothing on it: an outside ticket only shows up with logged time.
        self.assertEqual(len(build_dev_rows(ISSUES, "Bob", start, end, ["Alice", "Bob"], [support])), 2)

    def test_period_set_by_hand_replaces_the_jira_dates(self):
        period = SprintPeriod(
            start=datetime(2026, 3, 2, 0, 0, tzinfo=dt_timezone.utc), end=datetime(2026, 3, 4, 0, 0, tzinfo=dt_timezone.utc)
        )
        start, end = sprint_window(SPRINT, NOW, period)

        self.assertEqual((start, end), (period.start, period.end))
        self.assertEqual(logged_hours(ISSUES[0], "Alice", start, end), 3.0)

        content = LoggedWorkbookBuilder(sprint=SPRINT, issues=ISSUES, devs=["Alice"], period=period, now=NOW).build()
        rows = [[cell.value for cell in row] for row in load_workbook(BytesIO(content))["Period"].iter_rows()]
        self.assertEqual(rows[3], ["Bounds", "Set by hand"])

    def test_dev_tabs_and_unassigned_add_up_to_all_time_logged_in_the_sprint(self):
        start, end = sprint_window(SPRINT, NOW)
        dev_names = ["Alice", "Bob"]

        per_dev = sum(row["logged_hours"] for name in dev_names for row in build_dev_rows(ISSUES, name, start, end, dev_names))
        unassigned = sum(row["logged_hours"] for row in build_unassigned_rows(ISSUES, dev_names, start, end))

        self.assertEqual(per_dev + unassigned, 3.0 + 2.5 + 1.0 + 1.5 + 0.5)


class LoggedSheetViewTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="alice", password="secret")
        self.client.force_login(self.user)
        self.team = Team.objects.create(name="Core")
        TeamDev.objects.create(team=self.team, dev=Dev.objects.create(name="Alice"), order=0)

    def _fake_jira(self, team, worklogs=(), outside_issues=(), worklog_error=None):
        self.worklog_windows, self.issue_id_queries = [], []

        def get_worklogs_between(start, end):
            self.worklog_windows.append((start, end))
            if worklog_error:
                raise worklog_error
            return list(worklogs)

        def get_issues_by_ids(issue_ids):
            self.issue_id_queries.append(sorted(issue_ids))
            return [dict(issue) for issue in outside_issues if issue["id"] in issue_ids]

        return SimpleNamespace(
            cfg=SimpleNamespace(team=team, jira_project_id=7),
            get_sprint=lambda sprint_id: {**SPRINT, "url": "https://jira.example/issues/?jql=sprint%3D2446"},
            get_sprint_issues_with_worklogs=lambda sprint_id: ISSUES,
            get_worklogs_between=get_worklogs_between,
            get_issues_by_ids=get_issues_by_ids,
        )

    @patch("sprint.views.get_connected_jira_for_user")
    def test_logged_sheet_is_downloaded(self, mock_get_jira):
        mock_get_jira.return_value = self._fake_jira(self.team)

        response = self.client.get(reverse("sprint-logged-xls-export", args=[2446]))

        self.assertEqual(response.status_code, 200)
        self.assertIn('filename="sprint_2446_logged.xlsx"', response["Content-Disposition"])
        self.assertEqual(load_workbook(BytesIO(response.content)).sheetnames, ["Alice", "Unassigned", "Period"])
        # The period's worklogs are read once, between the sprint bounds.
        self.assertEqual(
            [(start.isoformat(), end.isoformat()) for start, end in self.worklog_windows],
            [("2026-03-02T08:00:00+01:00", "2026-03-13T18:00:00+01:00")],
        )
        self.assertEqual(self.issue_id_queries, [])

    @patch("sprint.views.get_connected_jira_for_user")
    def test_logged_sheet_adds_time_logged_outside_the_sprint(self, mock_get_jira):
        # Alice has no Jira ID: her worklogs are matched by name, like the sprint's.
        worklogs = [
            {"issue_id": "900", "author": ALICE, "started": "2026-03-04T14:00:00.000+0100", "seconds": 7200},
            # On a sprint ticket: already counted with the sprint, not fetched again.
            {"issue_id": "101", "author": ALICE, "started": "2026-03-04T15:00:00.000+0100", "seconds": 3600},
            # Someone outside the team: ignored.
            {"issue_id": "901", "author": CARL, "started": "2026-03-04T15:00:00.000+0100", "seconds": 3600},
        ]
        support = {**ISSUES[1], "id": "900", "key": "SUP-9", "summary": "Support", "worklogs": []}
        mock_get_jira.return_value = self._fake_jira(self.team, worklogs=worklogs, outside_issues=[support])

        response = self.client.get(reverse("sprint-logged-xls-export", args=[2446]))

        self.assertEqual(self.issue_id_queries, [["900"]])
        alice = [[cell.value for cell in row] for row in load_workbook(BytesIO(response.content))["Alice"].iter_rows()]
        self.assertEqual([row[0] for row in alice[1:3]], ["ABC-1 - Login page", "SUP-9 - Support"])
        self.assertEqual(alice[2][1:], [None, "Created during sprint", 2.0, "In Progress", "No"])

    @patch("sprint.views.get_connected_jira_for_user")
    def test_logged_sheet_survives_a_worklog_timeout(self, mock_get_jira):
        mock_get_jira.return_value = self._fake_jira(self.team, worklog_error=Timeout("read timed out"))

        response = self.client.get(reverse("sprint-logged-xls-export", args=[2446]))

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        period = [[cell.value for cell in row] for row in workbook["Period"].iter_rows()]
        self.assertEqual(period[-1][0], "Not fetched")
        self.assertIn("timeout", period[-1][1])

    @patch("sprint.views.get_connected_jira_for_user")
    def test_sprint_bounds_are_saved_and_reset_per_sprint(self, mock_get_jira):
        mock_get_jira.return_value = self._fake_jira(self.team)
        url = reverse("sprint-period", args=[2446])

        response = self.client.post(url, {"start": "2026-03-02T00:00", "end": "2026-03-13T23:59"})

        self.assertRedirects(response, reverse("sprint_xls_report", args=[2446]), fetch_redirect_response=False)
        period = SprintPeriod.objects.get(team=self.team, sprint_id=2446)
        self.assertEqual(timezone.localtime(period.start).strftime("%Y-%m-%d %H:%M"), "2026-03-02 00:00")
        self.assertFalse(SprintPeriod.objects.filter(sprint_id=2447).exists())

        page = self.client.get(reverse("sprint_xls_report", args=[2446]))
        self.assertContains(page, 'value="2026-03-02T00:00"')

        self.client.post(url, {"start": "2026-03-13T00:00", "end": "2026-03-02T00:00"})  # end before start
        self.assertEqual(SprintPeriod.objects.get(sprint_id=2446).end, period.end)

        self.client.post(url, {"action": "reset"})
        self.assertFalse(SprintPeriod.objects.filter(sprint_id=2446).exists())

    @patch("sprint.views.get_connected_jira_for_user")
    def test_logged_sheet_without_team_explains_what_is_missing(self, mock_get_jira):
        mock_get_jira.return_value = self._fake_jira(None)

        response = self.client.get(reverse("sprint-logged-xls-export", args=[2446]))

        self.assertEqual(response.status_code, 409)
        self.assertTemplateUsed(response, "jira/jira_unavailable.html")

    @patch("sprint.views.get_connected_jira_for_user")
    def test_capacity_page_shows_logged_sheet_button(self, mock_get_jira):
        mock_get_jira.return_value = self._fake_jira(self.team)
        self.user.language = "en"
        self.user.save(update_fields=["language"])

        response = self.client.get(reverse("sprint_xls_report", args=[2446]))

        self.assertContains(response, reverse("sprint-logged-xls-export", args=[2446]))
        self.assertContains(response, "Generate logged sheet")
        self.assertContains(response, 'data-field="jira_username"')
        self.assertContains(response, reverse("sprint-period", args=[2446]))

    @patch("sprint.views.get_connected_jira_for_user")
    def test_capacity_page_is_translated(self, mock_get_jira):
        mock_get_jira.return_value = self._fake_jira(self.team)
        expected = {
            "fr": ("Générer le rapport des temps prestés", "Enregistrer les bornes", "ID Jira"),
            "nl": ("Urenrapport genereren", "Grenzen opslaan", "Jira-ID"),
        }
        for language, texts in expected.items():
            with self.subTest(language=language):
                self.user.language = language
                self.user.save(update_fields=["language"])

                response = self.client.get(reverse("sprint_xls_report", args=[2446]))

                for text in texts:
                    self.assertContains(response, text)

    @patch("sprint.views.get_connected_jira_for_user")
    def test_open_in_jira_links_to_the_sprint(self, mock_get_jira):
        mock_get_jira.return_value = self._fake_jira(self.team)

        response = self.client.get(reverse("sprint_xls_report", args=[2446]))

        self.assertContains(response, 'href="https://jira.example/issues/?jql=sprint%3D2446"')

    def test_weighted_throughput_distribution_groups_done_by_story_points(self):
        jira = SimpleNamespace(
            JIRA_ONLY_STANDARD_TYPES="issuetype in standardIssueTypes()",
            get_cached_sprint_issues=lambda sprint_id, jql_extra=None, full=None: [
                {"state": "Done", "status_category": "done", "estimation": {"story_points": 3}},
                {"state": "Done", "status_category": "done", "estimation": {"story_points": 3}},
                {"state": "Done", "status_category": "done", "estimation": {"story_points": 5}},
                {"state": "Done", "status_category": "done", "estimation": {"story_points": None}},
                {"state": "Closed", "status_category": "done", "estimation": {"story_points": 8}},
                {"state": "In Progress", "status_category": "indeterminate", "estimation": {"story_points": 8}},
            ],
        )

        result = build_weighted_throughput_distribution(jira, 12)

        self.assertEqual(result["distribution"], {"labels": ["0", "3", "5", "8"], "counts": [1, 2, 1, 1]})
        self.assertEqual(result["done_count"], 5)
        self.assertEqual(result["estimated_done_count"], 5)

    @patch("sprint.views.get_connected_jira_for_user")
    def test_weighted_throughput_distribution_page_renders(self, mock_get_jira):
        mock_get_jira.return_value = SimpleNamespace(
            JIRA_ONLY_STANDARD_TYPES="issuetype in standardIssueTypes()",
            get_sprint=lambda sprint_id: {
                "id": sprint_id,
                "name": "Sprint 12",
                "state": "active",
                "startDate": "2026-03-01T08:00:00.000+0000",
                "endDate": "2026-03-14T18:00:00.000+0000",
            },
            get_url_jql=lambda jql: f"https://jira.example/issues/?jql={jql}",
            get_cached_sprint_issues=lambda sprint_id, jql_extra=None, full=None: [
                {"state": "Done", "estimation": {"story_points": 3}},
            ],
        )

        response = self.client.get(reverse("sprint-weighted-throughput-distribution", args=[12]))

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "sprint/sprint_weighted_throughput_distribution.html")
        self.assertContains(response, "Distribution du débit par pondération")


SCOPE_REPORT = {
    "completedIssues": [
        {"key": "BKL1-1", "summary": "Planned and done", "statusId": "10002", "assigneeName": "Martin, Alice"},
        {"key": "BKL1-2", "summary": "Added and done", "statusId": "10002"},
        {"key": "BKL1-3", "summary": "Planned and canceled", "statusId": "10009"},
    ],
    "issuesNotCompletedInCurrentSprint": [],
    "puntedIssues": [
        {"key": "BKL1-4", "summary": "Planned, carried over", "statusId": "3"},
        {"key": "BKL1-5", "summary": "Added, carried over", "statusId": "3"},
    ],
    "issuesCompletedInAnotherSprint": [{"key": "BKL1-6", "summary": "Done elsewhere", "statusId": "10002"}],
    "issueKeysAddedDuringSprint": {"BKL1-2": True, "BKL1-5": True},
    "entityData": {
        "statuses": {
            "10002": {"statusName": "Done", "status": {"statusCategory": {"key": "done"}}},
            "10009": {"statusName": "Canceled", "status": {"statusCategory": {"key": "done"}}},
            "3": {"statusName": "In Progress", "status": {"statusCategory": {"key": "indeterminate"}}},
        }
    },
}


class SprintScopeFlowTests(TestCase):
    def test_flows_split_planned_and_added_issues(self):
        scope = build_scope_flow(SCOPE_REPORT, lambda key: f"https://jira.example/browse/{key}")

        flows = {(flow["from"], flow["to"]): flow["flow"] for flow in scope["flows"]}
        self.assertEqual(
            flows,
            {
                ("origin:planned", "outcome:done"): 2,
                ("origin:added", "outcome:done"): 1,
                ("origin:planned", "outcome:removed"): 1,
                ("origin:added", "outcome:removed"): 1,
                ("origin:planned", "outcome:done_elsewhere"): 1,
                ("outcome:done", "end:Done"): 2,
                ("outcome:done", "end:Canceled"): 1,
                ("outcome:removed", "end:In Progress"): 2,
                ("outcome:done_elsewhere", "end:Done"): 1,
            },
        )
        self.assertEqual(scope["totals"], {"planned": 4, "added": 2})
        self.assertEqual(scope["status_categories"]["In Progress"], "indeterminate")
        self.assertFalse(scope["has_history"])
        first = scope["issues"][0]
        self.assertEqual((first["key"], first["origin"], first["assignee"]), ("BKL1-1", "planned", "Martin, Alice"))

    def test_statuses_at_sprint_start_and_end_come_from_the_changelog(self):
        start = datetime(2026, 3, 2, 8, tzinfo=dt_timezone.utc)
        end = datetime(2026, 3, 13, 18, tzinfo=dt_timezone.utc)
        histories = {
            # Planned, To Do at the start, done during the sprint.
            "BKL1-1": {"state": "Done", "created": "2026-02-01T10:00:00.000+0000", "status_changes": [
                {"at": "2026-03-05T10:00:00.000+0000", "from": "To Do", "to": "Done"},
            ]},
            # Planned, already In Progress at the start; reopened after the sprint end.
            "BKL1-4": {"state": "In Progress", "created": "2026-02-01T10:00:00.000+0000", "status_changes": [
                {"at": "2026-02-20T10:00:00.000+0000", "from": "To Do", "to": "In Progress"},
                {"at": "2026-03-10T10:00:00.000+0000", "from": "In Progress", "to": "In Test"},
                {"at": "2026-03-20T10:00:00.000+0000", "from": "In Test", "to": "In Progress"},
            ]},
            # Added during the sprint: no start status.
            "BKL1-5": {"state": "In Progress", "created": "2026-03-04T10:00:00.000+0000", "status_changes": []},
        }

        scope = build_scope_flow(
            SCOPE_REPORT, lambda key: key, histories=histories, start=start, end=end,
            status_categories={"To Do": "new", "In Test": "indeterminate"},
        )

        issues = {issue["key"]: issue for issue in scope["issues"]}
        self.assertEqual((issues["BKL1-1"]["status_start"], issues["BKL1-1"]["status_end"]), ("To Do", "Done"))
        self.assertEqual(
            (issues["BKL1-4"]["status_start"], issues["BKL1-4"]["status_end"], issues["BKL1-4"]["status_current"]),
            ("In Progress", "In Test", "In Progress"),
        )
        self.assertIsNone(issues["BKL1-5"]["status_start"])
        # No history for BKL1-3: its current status stands in for the end status.
        self.assertEqual(issues["BKL1-3"]["status_end"], "Canceled")
        flows = {(flow["from"], flow["to"]): flow["flow"] for flow in scope["flows"]}
        self.assertEqual(flows[("start:To Do", "origin:planned")], 1)
        self.assertEqual(flows[("start:In Progress", "origin:planned")], 1)
        self.assertEqual(flows[("outcome:removed", "end:In Test")], 1)
        self.assertEqual(scope["columns"]["start:To Do"], 0)
        self.assertEqual(scope["columns"]["end:In Test"], 3)
        self.assertTrue(scope["has_history"])

    def _login_and_fake_jira(self, mock_get_jira, **overrides):
        user = get_user_model().objects.create_user(username="alice", password="secret")
        self.client.force_login(user)
        fake = dict(
            cfg=SimpleNamespace(jira_board_id=638),
            get_sprint=lambda sprint_id: {**SPRINT, "name": "Sprint 12", "url": "https://jira.example/"},
            get_sprint_scope_report=lambda sprint_id: SCOPE_REPORT,
            get_issue_url=lambda key: f"https://jira.example/browse/{key}",
            get_status_histories=lambda keys: {},
            status_category_keys=lambda: {},
        )
        fake.update(overrides)
        mock_get_jira.return_value = SimpleNamespace(**fake)
        return user

    @patch("sprint.views.get_connected_jira_for_user")
    def test_scope_page_renders(self, mock_get_jira):
        requested = []
        user = self._login_and_fake_jira(
            mock_get_jira, get_status_histories=lambda keys: requested.extend(keys) or {}
        )

        response = self.client.get(reverse("sprint-scope-flow", args=[2446]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "chartjs-chart-sankey")
        self.assertContains(response, "BKL1-5")
        self.assertEqual(requested, ["BKL1-1", "BKL1-2", "BKL1-3", "BKL1-4", "BKL1-5", "BKL1-6"])

        user.language = "fr"
        user.save(update_fields=["language"])
        response = self.client.get(reverse("sprint-scope-flow", args=[2446]))
        self.assertContains(response, "Périmètre du sprint")
        self.assertContains(response, "Retirés / reportés")

    @patch("sprint.views.get_connected_jira_for_user")
    def test_scope_page_falls_back_to_current_statuses_when_history_times_out(self, mock_get_jira):
        def timeout(keys):
            raise Timeout("read timed out")

        self._login_and_fake_jira(mock_get_jira, get_status_histories=timeout)

        response = self.client.get(reverse("sprint-scope-flow", args=[2446]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "alert-warning")
        self.assertFalse(response.context["scope"]["has_history"])
