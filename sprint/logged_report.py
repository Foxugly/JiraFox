"""
"Logged sheet": one tab per team developer listing, for each ticket the developer worked on (time
logged during the sprint) or is assigned to in the sprint:

    Ticket | Planned (h) | Status at sprint start | Logged (h) | Status at sprint end | In sprint

Tickets outside the sprint the developer logged time on come last ("In sprint" = No, no planned
time), followed by a Total row and an "Outside sprint" subtotal. An "Unassigned" tab lists the
time logged on sprint tickets by people matching no team developer (one row per ticket and author),
so name mismatches between the developers and Jira are visible instead of silently dropped. A last
"Period" tab states the bounds used, which come from Jira unless set by hand (SprintPeriod).
"""
from collections import defaultdict
from dataclasses import dataclass
import re
from datetime import datetime
from io import BytesIO
from typing import Iterable

from dateutil import parser as date_parser
from django.utils import timezone
from jira.exceptions import JIRAError
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from requests import RequestException

from jiramodule.exceptions import JiraUnavailable, describe_jira_error
from jiramodule.services.jira_client import JiraManager
from team.models import TeamDev

from .models import SprintPeriod

HEADERS = ("Ticket", "Planned (h)", "Status at sprint start", "Logged (h)", "Status at sprint end", "In sprint")
UNASSIGNED_HEADERS = HEADERS[:5] + ("Logged by",)
UNASSIGNED_SHEET_TITLE = "Unassigned"
PERIOD_SHEET_TITLE = "Period"
CREATED_DURING_SPRINT = "Created during sprint"
SHEET_TITLE_MAX_LENGTH = 31
INVALID_SHEET_TITLE_CHARS = re.compile(r"[\[\]:*?/\\]")


@dataclass(frozen=True)
class DevIdentity:
    """A team developer as the logged sheet matches them against Jira people."""

    name: str
    jira_username: str = ""

    @classmethod
    def of(cls, value: "DevIdentity | str") -> "DevIdentity":
        return value if isinstance(value, cls) else cls(name=value)

    @classmethod
    def from_dev(cls, dev) -> "DevIdentity":
        return cls(name=dev.name, jira_username=(dev.jira_username or "").strip())


def parse_jira_datetime(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        parsed = date_parser.isoparse(value)
    except (TypeError, ValueError):
        return None
    return timezone.make_aware(parsed) if timezone.is_naive(parsed) else parsed


def _username(person: dict) -> str:
    return (person.get("name") or "").strip().casefold()


def _exact_names(person: dict) -> set[str]:
    names = {(person.get("displayName") or "").strip(), (person.get("name") or "").strip()}
    return {name.casefold() for name in names if name}


def _first_name(person: dict) -> str | None:
    display_name = (person.get("displayName") or "").strip()
    return display_name.split(", ", 1)[1].casefold() if ", " in display_name else None


def person_matches_dev(person: dict | None, dev: DevIdentity | str) -> bool:
    """
    A developer with a Jira username matches that username only. Otherwise, same rule as the
    capacity sheet: Jira display names look like "Last, First" and the developer name holds the
    first name; the full display name and the Jira username are accepted too.
    """
    dev = DevIdentity.of(dev)
    if not person or not dev.name:
        return False
    if dev.jira_username:
        return _username(person) == dev.jira_username.casefold()
    target = dev.name.strip().casefold()
    return target in _exact_names(person) or target == _first_name(person)


def resolve_dev(person: dict | None, devs: Iterable[DevIdentity | str]) -> DevIdentity | None:
    """
    The one team developer a Jira person is, so that no time is counted twice. A Jira username wins,
    then an exact name match, then a first-name match; a person fitting several developers at the
    same level is ambiguous and left unassigned, where it shows up instead of being doubled.
    """
    if not person:
        return None
    devs = [DevIdentity.of(dev) for dev in devs]
    username = _username(person)
    by_username = [dev for dev in devs if dev.jira_username and dev.jira_username.casefold() == username]
    by_name = [dev for dev in devs if not dev.jira_username]
    exact = _exact_names(person)
    first_name = _first_name(person)
    for matches in (
        by_username,
        [dev for dev in by_name if dev.name.strip().casefold() in exact],
        [dev for dev in by_name if first_name and dev.name.strip().casefold() == first_name],
    ):
        if matches:
            return matches[0] if len(matches) == 1 else None
    return None


def status_at(issue: dict, moment: datetime) -> str | None:
    """Status of the issue at `moment`, rebuilt from its changelog. None if it did not exist yet."""
    created = parse_jira_datetime(issue.get("created"))
    if created and created > moment:
        return None

    changes = sorted(
        (
            (when, change.get("from"), change.get("to"))
            for change in issue.get("status_changes") or []
            if (when := parse_jira_datetime(change.get("at")))
        ),
        key=lambda change: change[0],
    )
    passed = [change for change in changes if change[0] <= moment]
    if passed:
        return passed[-1][2]
    if changes:
        return changes[0][1]
    return issue.get("state")


def sprint_window(
    sprint: dict, now: datetime, period: SprintPeriod | None = None
) -> tuple[datetime | None, datetime]:
    """(start, end) of the sprint: the hand-set period if any, else Jira's dates; never past now."""
    if period:
        start, end = period.start, period.end
    else:
        start = parse_jira_datetime(sprint.get("startDate"))
        end = parse_jira_datetime(sprint.get("completeDate")) or parse_jira_datetime(sprint.get("endDate"))
    return start, min(end, now) if end else now


def sprint_worklogs(issue: dict, start: datetime | None, end: datetime) -> list[dict]:
    if start is None:
        # A sprint that never started has no logged time; without a lower bound the whole
        # history of its (carried-over) tickets would be counted.
        return []
    worklogs = []
    for worklog in issue.get("worklogs") or []:
        started = parse_jira_datetime(worklog.get("started"))
        if started is None or started < start or started > end:
            continue
        worklogs.append(worklog)
    return worklogs


def worklog_seconds(worklog: dict) -> float:
    seconds = worklog.get("seconds")
    return float(seconds) if seconds is not None else float(worklog.get("hours") or 0) * 3600


def seconds_to_hours(seconds: float) -> float:
    return round(seconds / 3600, 2)


def logged_hours(
    issue: dict,
    dev: DevIdentity | str,
    start: datetime | None,
    end: datetime,
    devs: Iterable[DevIdentity | str] | None = None,
) -> float:
    dev = DevIdentity.of(dev)
    devs = list(devs) if devs is not None else [dev]
    total = sum(
        worklog_seconds(worklog)
        for worklog in sprint_worklogs(issue, start, end)
        if resolve_dev(worklog.get("author"), devs) == dev
    )
    return seconds_to_hours(total)


def issue_row(issue: dict, hours: float, start: datetime | None, end: datetime, in_sprint: bool = True) -> dict:
    start_status = status_at(issue, start) if start else issue.get("state")
    return {
        "ticket": f"{issue.get('key', '')} - {issue.get('summary', '')}",
        "url": issue.get("url") or "",
        # Planned time is the sprint's: a ticket outside the sprint was not planned in it.
        "planned_hours": float(issue.get("original_estimate_hours") or 0) if in_sprint else None,
        "start_status": start_status or CREATED_DURING_SPRINT,
        "logged_hours": hours,
        "end_status": status_at(issue, end) or "",
        "in_sprint": in_sprint,
    }


def build_dev_rows(
    issues: Iterable[dict],
    dev: DevIdentity | str,
    start: datetime | None,
    end: datetime,
    devs: Iterable[DevIdentity | str] | None = None,
    outside_issues: Iterable[dict] = (),
) -> list[dict]:
    dev = DevIdentity.of(dev)
    devs = list(devs) if devs is not None else [dev]
    rows = []
    for issue in issues:
        hours = logged_hours(issue, dev, start, end, devs)
        if hours <= 0 and not person_matches_dev(issue.get("assignee"), dev):
            continue
        rows.append(issue_row(issue, hours, start, end))
    for issue in outside_issues:
        hours = logged_hours(issue, dev, start, end, devs)
        if hours > 0:
            rows.append(issue_row(issue, hours, start, end, in_sprint=False))
    return rows


def author_label(author: dict | None) -> str:
    if not author:
        return "Unknown"
    display_name = author.get("displayName") or ""
    username = author.get("name") or ""
    if display_name and username:
        return f"{display_name} ({username})"
    return display_name or username or "Unknown"


def build_unassigned_rows(
    issues: Iterable[dict], devs: Iterable[DevIdentity | str], start: datetime | None, end: datetime
) -> list[dict]:
    """Time logged during the sprint by authors matching no team developer, or several of them."""
    devs = list(devs)
    rows = []
    for issue in issues:
        seconds_by_author: dict[str, float] = defaultdict(float)
        for worklog in sprint_worklogs(issue, start, end):
            author = worklog.get("author")
            if resolve_dev(author, devs):
                continue
            seconds_by_author[author_label(author)] += worklog_seconds(worklog)
        for index, (author, seconds) in enumerate(sorted(seconds_by_author.items())):
            row = {**issue_row(issue, seconds_to_hours(seconds), start, end), "author": author}
            if index:
                # Planned time belongs to the ticket: count it once, not once per author.
                row["planned_hours"] = None
            rows.append(row)
    return rows


def sheet_title(name: str, used: set[str]) -> str:
    base = INVALID_SHEET_TITLE_CHARS.sub("-", name).strip().strip("'") or "Dev"
    base = base[:SHEET_TITLE_MAX_LENGTH]
    title, suffix = base, 2
    while title.casefold() in used:
        tail = f" ({suffix})"
        title = f"{base[:SHEET_TITLE_MAX_LENGTH - len(tail)]}{tail}"
        suffix += 1
    used.add(title.casefold())
    return title


def _excel_datetime(value: datetime | None) -> datetime | None:
    # Excel has no time zones: show the local time.
    return timezone.localtime(value).replace(tzinfo=None) if value else None


class LoggedWorkbookBuilder:
    header_fill = PatternFill(start_color="ff6565", end_color="ff6565", fill_type="solid")
    total_fill = PatternFill(start_color="ff8f8f", end_color="ff8f8f", fill_type="solid")
    border = Border(*(Side(style="thin"),) * 4)
    column_widths = (70, 13, 26, 13, 26, 34)

    def __init__(
        self,
        *,
        sprint: dict,
        issues: list[dict],
        devs: list[DevIdentity | str],
        outside_issues: list[dict] = (),
        period: SprintPeriod | None = None,
        now: datetime | None = None,
        warnings: list[str] = (),
    ):
        self.sprint = sprint
        self.warnings = list(warnings)
        self.issues = issues
        self.devs = [DevIdentity.of(dev) for dev in devs]
        self.outside_issues = list(outside_issues)
        self.period = period
        self.start, self.end = sprint_window(sprint, now or timezone.now(), period)
        self.workbook = Workbook()

    def build(self) -> bytes:
        self.workbook.remove(self.workbook.active)
        used_titles: set[str] = set()
        if not self.devs:
            worksheet = self.workbook.create_sheet("No developers")
            worksheet["A1"] = "The team has no developer: add them on the capacity sheet page."
        for dev in self.devs:
            worksheet = self.workbook.create_sheet(sheet_title(dev.name, used_titles))
            rows = build_dev_rows(self.issues, dev, self.start, self.end, self.devs, self.outside_issues)
            self._populate(worksheet, rows)

        worksheet = self.workbook.create_sheet(sheet_title(UNASSIGNED_SHEET_TITLE, used_titles))
        self._populate(
            worksheet,
            build_unassigned_rows(self.issues, self.devs, self.start, self.end),
            headers=UNASSIGNED_HEADERS,
        )
        self._populate_period(self.workbook.create_sheet(sheet_title(PERIOD_SHEET_TITLE, used_titles)))

        buffer = BytesIO()
        self.workbook.save(buffer)
        return buffer.getvalue()

    def _row_values(self, row: dict, headers: tuple[str, ...]) -> tuple:
        values = (row["ticket"], row["planned_hours"], row["start_status"], row["logged_hours"], row["end_status"])
        if headers == UNASSIGNED_HEADERS:
            return values + (row["author"],)
        return values + ("Yes" if row["in_sprint"] else "No",)

    def _populate(self, worksheet, rows: list[dict], headers: tuple[str, ...] = HEADERS) -> None:
        for column, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=1, column=column, value=header)
            cell.font = Font(bold=True)
            cell.fill = self.header_fill
            cell.border = self.border
        worksheet.freeze_panes = "A2"

        for row_index, row in enumerate(rows, start=2):
            for column, value in enumerate(self._row_values(row, headers), start=1):
                worksheet.cell(row=row_index, column=column, value=value).border = self.border
            if row["url"]:
                worksheet.cell(row=row_index, column=1).hyperlink = row["url"]

        last_data_row = len(rows) + 1
        total_rows = [
            {
                1: "Total",
                2: f"=SUM(B2:B{last_data_row})" if rows else 0,
                4: f"=SUM(D2:D{last_data_row})" if rows else 0,
            }
        ]
        if headers == HEADERS:
            total_rows.append(
                {1: "Outside sprint", 4: f'=SUMIF(F2:F{last_data_row},"No",D2:D{last_data_row})' if rows else 0}
            )
        for offset, totals in enumerate(total_rows):
            for column in range(1, len(headers) + 1):
                cell = worksheet.cell(row=last_data_row + 1 + offset, column=column, value=totals.get(column))
                cell.font = Font(bold=True)
                cell.fill = self.total_fill
                cell.border = self.border

        for column, width in enumerate(self.column_widths[: len(headers)], start=1):
            worksheet.column_dimensions[chr(ord("A") + column - 1)].width = width

    def _populate_period(self, worksheet) -> None:
        lines = (
            ("Sprint", self.sprint.get("name") or self.sprint.get("id")),
            ("Start", _excel_datetime(self.start)),
            ("End", _excel_datetime(self.end)),
            ("Bounds", "Set by hand" if self.period else "Jira sprint dates"),
            *(("Not fetched", f"Time logged outside the sprint: {warning}") for warning in self.warnings),
        )
        for row_index, (label, value) in enumerate(lines, start=1):
            worksheet.cell(row=row_index, column=1, value=label).font = Font(bold=True)
            cell = worksheet.cell(row=row_index, column=2, value=value)
            if isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
        worksheet.column_dimensions["A"].width = 12
        worksheet.column_dimensions["B"].width = 30


def fetch_outside_issues(
    jira: JiraManager, issues: list[dict], devs: list[DevIdentity], start: datetime, end: datetime
) -> list[dict]:
    """
    Issues outside the sprint on which a team developer logged time during the period, carrying
    those worklogs only. Found from the worklogs of the period, matched like the sprint's.
    """
    sprint_ids = {issue.get("id") for issue in issues}
    worklogs_by_issue: dict[str, list[dict]] = defaultdict(list)
    for worklog in jira.get_worklogs_between(start, end):
        if worklog["issue_id"] not in sprint_ids and resolve_dev(worklog.get("author"), devs):
            worklogs_by_issue[worklog["issue_id"]].append(worklog)
    if not worklogs_by_issue:
        return []
    outside = jira.get_issues_by_ids(list(worklogs_by_issue))
    for issue in outside:
        issue["worklogs"] = worklogs_by_issue.get(issue.get("id"), [])
    return outside


def create_logged_xlsx_bytes(jira: JiraManager, sprint_id: int, now: datetime | None = None) -> bytes:
    team = jira.cfg.team
    devs = [
        DevIdentity.from_dev(teamdev.dev)
        for teamdev in TeamDev.objects.filter(team=team).select_related("dev").order_by("order", "id")
    ]
    sprint = jira.get_sprint(sprint_id)
    period = SprintPeriod.objects.filter(team=team, sprint_id=sprint_id).first()
    issues = jira.get_sprint_issues_with_worklogs(sprint_id)

    start, end = sprint_window(sprint, now or timezone.now(), period)
    outside_issues: list[dict] = []
    warnings: list[str] = []
    if start:
        try:
            outside_issues = fetch_outside_issues(jira, issues, devs, start, end)
        except (JIRAError, RequestException, JiraUnavailable) as exc:
            # The sprint itself is still worth reporting: say what is missing instead.
            warnings.append(describe_jira_error(exc))

    return LoggedWorkbookBuilder(
        sprint=sprint,
        issues=issues,
        devs=devs,
        outside_issues=outside_issues,
        period=period,
        now=now,
        warnings=warnings,
    ).build()
